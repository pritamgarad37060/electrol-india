import os, sys, re, json
import openpyxl

CORP_FILE = 'scratch_corp.xlsx'
COUNCIL_FILE = 'scratch_council.xlsx'

CORP_CLEAN_MAP = {
    'Brihanmumbai Municipal Corporation (BMC)': 'Brihanmumbai (BMC)',
    'Pune Municipal Corporation (PMC)': 'Pune (PMC)',
    'Nagpur Municipal Corporation (NMC)': 'Nagpur (NMC)',
    'Thane Municipal Corporation (TMC)': 'Thane (TMC)',
    'Pimpari Chinchwad Municipal Corporation (PCMC)': 'Pimpri-Chinchwad (PCMC)',
    'Nashik Municipal Corporation (NMC)': 'Nashik (NMC)',
    'Navi Mumbai Municipal Corporation (NMMC)': 'Navi Mumbai (NMMC)',
    'Kalyan Dombivali Municipal Corporation (KDMC)': 'Kalyan-Dombivli (KDMC)',
    'Vasai Virar City Municipal Corporation (VVCMC)': 'Vasai-Virar (VVCMC)',
    'Chhatrapati Sambhajinagar Municipal': 'Chhatrapati Sambhajinagar',
    'Ahilyanagar Municipal Corporation': 'Ahilyanagar',
    'Solapur Municipal Corporation (SMC)': 'Solapur (SMC)',
    'Amravati Municipal Corporation (AMC)': 'Amravati (AMC)',
    'Kolhapur Municipal Corporation': 'Kolhapur',
    'Akola Municipal Corporation': 'Akola',
    'Bhiwandi-Nizampur Municipal Corporation (BNMC)': 'Bhiwandi-Nizampur (BNMC)',
    'Mira-Bhaindar Municipal Corporation (MBMC)': 'Mira-Bhayandar (MBMC)',
    'Ulhasnagar Municipal Corporation (UMC)': 'Ulhasnagar (UMC)',
    'Sangali-Miraj-Kupwad Municipal Corporation (SMKMC)': 'Sangli-Miraj-Kupwad',
    'Panvel Municipal Corporation': 'Panvel',
    'Jalgaon': 'Jalgaon',
    'Dhule Municipal Corporation (DMC)': 'Dhule',
    'Latur Municipal Corporation (LMC)': 'Latur',
    'Chandrapur Municipal Corporation (CMC)': 'Chandrapur',
    'Parbhani Municipal Corporation': 'Parbhani',
    'Nanded- Waghala Municipal Corporation (NWMC)': 'Nanded-Waghala',
    'Jalna Municipal Corporation': 'Jalna',
    'Malegaon Municipal Corporation (MNC)': 'Malegaon',
    'Ichalkaranji Municipal Corporation (IMC)': 'Ichalkaranji'
}

CORP_DISTRICT_MAP = {
    'Brihanmumbai (BMC)': 'Mumbai & Mumbai Suburban',
    'Pune (PMC)': 'Pune',
    'Nagpur (NMC)': 'Nagpur',
    'Thane (TMC)': 'Thane',
    'Pimpri-Chinchwad (PCMC)': 'Pune',
    'Nashik (NMC)': 'Nashik',
    'Navi Mumbai (NMMC)': 'Thane',
    'Kalyan-Dombivli (KDMC)': 'Thane',
    'Vasai-Virar (VVCMC)': 'Palghar',
    'Chhatrapati Sambhajinagar': 'Chhatrapati Sambhajinagar',
    'Ahilyanagar': 'Ahilyanagar',
    'Solapur (SMC)': 'Solapur',
    'Amravati (AMC)': 'Amravati',
    'Kolhapur': 'Kolhapur',
    'Akola': 'Akola',
    'Bhiwandi-Nizampur (BNMC)': 'Thane',
    'Mira-Bhayandar (MBMC)': 'Thane',
    'Ulhasnagar (UMC)': 'Thane',
    'Sangli-Miraj-Kupwad': 'Sangli',
    'Panvel': 'Raigad',
    'Jalgaon': 'Jalgaon',
    'Dhule': 'Dhule',
    'Latur': 'Latur',
    'Chandrapur': 'Chandrapur',
    'Parbhani': 'Parbhani',
    'Nanded-Waghala': 'Nanded',
    'Jalna': 'Jalna',
    'Malegaon': 'Nashik',
    'Ichalkaranji': 'Kolhapur'
}

def clean_corp_name(raw):
    raw_str = str(raw or '').strip()
    for k, v in CORP_CLEAN_MAP.items():
        if k.lower() in raw_str.lower():
            return v
    c = re.sub(r'[\uFFFD\?]', ' ', raw_str)
    c = re.sub(r'2026|All Wards|\(.*\)', '', c).strip()
    c = re.sub(r'Municipal Corporation.*', '', c).strip()
    return c or raw_str

def normalize_party(p):
    if not p or str(p).strip() == '':
        return 'Independent'
    p_str = str(p).strip()
    p_up = p_str.upper()

    if 'UBT' in p_up or 'UDDHAV' in p_up:
        return 'Shiv Sena (UBT)'
    if p_up in ['SHIVSENA', 'SHIVESENA', 'SHIV SENA', 'SHS'] or p_up.startswith('SHIV SENA'):
        return 'Shiv Sena'
    if 'BHARATIYA JANATA' in p_up or 'BHARTIYA JANATA' in p_up or p_up == 'BJP':
        return 'BJP'
    if 'SHARAD' in p_up or '(SP)' in p_up or 'SHARADCHANDRA' in p_up:
        return 'NCP (SP)'
    if 'NATIONALIST CONGRESS' in p_up or p_up == 'NCP':
        return 'NCP'
    if 'INDIAN NATIONAL CONGRESS' in p_up or p_up == 'INC' or 'CONGRESS' in p_up:
        return 'INC'
    if 'MAHARASHTRA NAVNIRMAN' in p_up or p_up == 'MNS':
        return 'MNS'
    if 'ALL INDIA MAJLIS' in p_up or 'AIMIM' in p_up:
        return 'AIMIM'
    if 'VANCHIT' in p_up or 'VBA' in p_up:
        return 'VBA'
    if 'BAHUJAN VIKAS' in p_up or 'BVA' in p_up:
        return 'BVA'
    if 'PEASANTS' in p_up or 'WORKERS' in p_up or 'PWPI' in p_up:
        return 'PWPI'
    if 'SAMAJWADI' in p_up or p_up == 'SP':
        return 'SP'
    if 'BAHUJAN SAMAJ' in p_up or p_up == 'BSP':
        return 'BSP'
    if 'COMMUNIST' in p_up or 'CPI' in p_up:
        return 'CPI(M)'
    if 'REPUBLICAN' in p_up or 'RPI' in p_up:
        return 'RPI(A)'
    if 'NOTA' in p_up or 'NONE OF' in p_up:
        return 'NOTA'
    if 'INDEPENDENT' in p_up or 'IND' in p_up:
        return 'Independent'
    return p_str

print("Loading Corporation workbook...")
wb_corp = openpyxl.load_workbook(CORP_FILE, data_only=True)

# 1. Reservations
res_sheet = wb_corp['Ward_Reservations_2026']
reservations_2026 = {}
for i, r in enumerate(res_sheet.iter_rows(values_only=True)):
    if i == 0: continue
    c_name = str(r[1] or '').strip()
    w_no = r[2]
    res_val = str(r[3] or '').strip()
    reservations_2026[(c_name, str(w_no))] = res_val

# 2. Corp metadata
corp_meta_sheet = wb_corp['Corporations']
corps_meta = []
for i, r in enumerate(corp_meta_sheet.iter_rows(values_only=True)):
    if i == 0:
        meta_cols = [str(x) for x in r]
    else:
        corps_meta.append(dict(zip(meta_cols, r)))

# 3. Candidates 2026
c26_sheet = wb_corp['Candidates_2026']
c26_data = {}
for i, r in enumerate(c26_sheet.iter_rows(values_only=True)):
    if i == 0:
        c26_cols = [str(x) for x in r]
    else:
        corp = str(r[1] or '').strip()
        if corp not in c26_data:
            c26_data[corp] = []
        c26_data[corp].append(dict(zip(c26_cols, r)))

# 4. Candidates 2017
c17_sheet = wb_corp['Candidates_2017']
c17_data = {}
for i, r in enumerate(c17_sheet.iter_rows(values_only=True)):
    if i == 0:
        c17_cols = [str(x) for x in r]
    else:
        corp = str(r[1] or '').strip()
        if corp not in c17_data:
            c17_data[corp] = []
        c17_data[corp].append(dict(zip(c17_cols, r)))

print("Processing 29 Corporations...")
corps_list = []
corp_tally_2026 = {}
corp_tally_2017 = {}

def process_cands_for_wards(cands_list, corp_raw, is_2026=True):
    grouped = {}
    for row in cands_list:
        w_no = row['Ward']
        s_id = str(row.get('Seat', '') or '').strip()
        if s_id == str(w_no):
            s_id = ''
        key = (str(w_no), s_id)
        if key not in grouped:
            grouped[key] = []
        grouped[key].append(row)

    wards = []
    tally = {}
    for (w_no, s_id), cand_rows in grouped.items():
        # Sort candidates descending by Votes
        cand_rows.sort(key=lambda x: (x.get('Votes') or 0), reverse=True)
        winner = cand_rows[0]
        runner = cand_rows[1] if len(cand_rows) > 1 else None
        third = cand_rows[2] if len(cand_rows) > 2 else None
        fourth = cand_rows[3] if len(cand_rows) > 3 else None

        w_votes = int(winner['Votes'] or 0)
        r_votes = int(runner['Votes'] or 0) if runner else 0
        margin = w_votes - r_votes if runner else w_votes

        p_norm = normalize_party(winner.get('Political_Party'))
        tally[p_norm] = tally.get(p_norm, 0) + 1

        res_val = reservations_2026.get((corp_raw, w_no), '')
        if not res_val:
            res_val = str(winner.get('Ward_Reservation') or '')

        # Compact all candidates list
        compact_cands = []
        for c in cand_rows:
            compact_cands.append([
                str(c.get('Candidate', '')),
                normalize_party(c.get('Political_Party')),
                int(c.get('Votes') or 0),
                round(float(c.get('Vote_Share_Pct') or 0.0), 2)
            ])

        ward_obj = {
            'ward': int(w_no) if str(w_no).isdigit() else str(w_no),
            'seat': s_id,
            'ward_label': f"{w_no}{s_id}" if s_id else str(w_no),
            'candidate': str(winner.get('Candidate', '')),
            'party': p_norm,
            'votes': w_votes,
            'pct': round(float(winner.get('Vote_Share_Pct') or 0.0), 2),
            'res': res_val,
            'runner': str(runner.get('Candidate', '')) if runner else '',
            'runner_party': normalize_party(runner.get('Political_Party')) if runner else '',
            'runner_votes': r_votes,
            'runner_pct': round(float(runner.get('Vote_Share_Pct') or 0.0), 2) if runner else 0.0,
            'third': str(third.get('Candidate', '')) if third else '',
            'third_party': normalize_party(third.get('Political_Party')) if third else '',
            'third_votes': int(third['Votes'] or 0) if third else 0,
            'third_pct': round(float(third.get('Vote_Share_Pct') or 0.0), 2) if third else 0.0,
            'fourth': str(fourth.get('Candidate', '')) if fourth else '',
            'fourth_party': normalize_party(fourth.get('Political_Party')) if fourth else '',
            'fourth_votes': int(fourth['Votes'] or 0) if fourth else 0,
            'fourth_pct': round(float(fourth.get('Vote_Share_Pct') or 0.0), 2) if fourth else 0.0,
            'margin': margin,
            'cands': compact_cands
        }
        wards.append(ward_obj)

    # Sort wards by ward number
    def w_sort(x):
        w = x['ward']
        s = x['seat']
        return (int(w) if str(w).isdigit() else 9999, str(w), str(s))
    wards.sort(key=w_sort)
    return wards, tally

for r in corps_meta:
    raw_name = str(r.get('Municipal_Corporation') or r.get('Corporation') or '')
    cname = clean_corp_name(raw_name)
    dist = CORP_DISTRICT_MAP.get(cname, 'Maharashtra')

    c26_sub = c26_data.get(raw_name, [])
    if not c26_sub:
        for k in c26_data.keys():
            if cname.lower() in k.lower():
                c26_sub = c26_data[k]
                break

    c17_sub = c17_data.get(raw_name, [])
    if not c17_sub:
        for k in c17_data.keys():
            if cname.lower() in k.lower():
                c17_sub = c17_data[k]
                break

    wards_26, tally_26 = process_cands_for_wards(c26_sub, raw_name, is_2026=True)
    wards_17, tally_17 = process_cands_for_wards(c17_sub, raw_name, is_2026=False)

    for p, cnt in tally_26.items():
        corp_tally_2026[p] = corp_tally_2026.get(p, 0) + cnt
    for p, cnt in tally_17.items():
        corp_tally_2017[p] = corp_tally_2017.get(p, 0) + cnt

    slug = 'bmc' if 'brihanmumbai' in cname.lower() or 'bmc' in cname.lower() else re.sub(r'[^a-z0-9]+', '-', cname.lower()).strip('-')

    corp_obj = {
        'id': len(corps_list) + 1,
        'slug': slug,
        'name': cname,
        'raw_name': raw_name,
        'district': dist,
        'total_pop': int(r['Total_Population']) if r.get('Total_Population') else None,
        'sc_pop': int(r['SC_Population']) if r.get('SC_Population') else None,
        'st_pop': int(r['ST_Population']) if r.get('ST_Population') else None,
        'ward_count_2026': len(set(x['ward'] for x in wards_26)),
        'seat_count_2026': len(wards_26),
        'party_tally_2026': tally_26,
        'winners_2026': wards_26,
        'ward_count_2017': len(set(x['ward'] for x in wards_17)),
        'seat_count_2017': len(wards_17),
        'party_tally_2017': tally_17,
        'winners_2017': wards_17
    }
    corps_list.append(corp_obj)

print("Loading Municipal Councils workbook...")
wb_counc = openpyxl.load_workbook(COUNCIL_FILE, data_only=True)

# 1. Councils meta
counc_meta_sheet = wb_counc['Municipal_Councils']
counc_meta = []
for i, r in enumerate(counc_meta_sheet.iter_rows(values_only=True)):
    if i == 0:
        meta_cols = [str(x) for x in r]
    else:
        counc_meta.append(dict(zip(meta_cols, r)))

# 2. Candidate results
counc_cand_sheet = wb_counc['Candidate_Results']
counc_data = {}
for i, r in enumerate(counc_cand_sheet.iter_rows(values_only=True)):
    if i == 0:
        cand_cols = [str(x) for x in r]
    else:
        c_name = str(r[2] or '').strip()
        if c_name not in counc_data:
            counc_data[c_name] = []
        counc_data[c_name].append(dict(zip(cand_cols, r)))

print("Processing 224 Municipal Councils with ALL candidates (1st, 2nd, 3rd, 4th)...")
councils_list = []
council_tally_2025 = {}
council_pres_tally = {}

def clean_council_name(name):
    c = re.sub(r'Municipal Council\s*', '', str(name or ''), flags=re.IGNORECASE).strip()
    c = re.sub(r'[\uFFFD\?]', ' ', c).strip()
    return c

for r in counc_meta:
    raw_name = str(r['Council'])
    cname = clean_council_name(raw_name)

    c_rows = counc_data.get(raw_name, [])
    if not c_rows:
        for k in counc_data.keys():
            if cname.lower() in k.lower():
                c_rows = counc_data[k]
                break

    # President candidate extraction (ALL CANDIDATES: Winner, Runner-up, 3rd, 4th, NOTA)
    pres_rows = [x for x in c_rows if str(x.get('Post_Type', '')).lower() == 'president']
    pres_rows.sort(key=lambda x: (x.get('Votes') or 0), reverse=True)
    president_obj = None
    if pres_rows:
        pw = pres_rows[0]
        prunner = pres_rows[1] if len(pres_rows) > 1 else None
        pthird = pres_rows[2] if len(pres_rows) > 2 else None
        pfourth = pres_rows[3] if len(pres_rows) > 3 else None

        pw_votes = int(pw.get('Votes') or 0)
        pr_votes = int(prunner.get('Votes') or 0) if prunner else 0
        pmargin = pw_votes - pr_votes if prunner else pw_votes

        p_norm = normalize_party(pw.get('Party'))
        
        pres_cands = []
        for p in pres_rows:
            pres_cands.append([
                str(p.get('Candidate', '')),
                normalize_party(p.get('Party')),
                int(p.get('Votes') or 0),
                round(float(p.get('Vote_Share_Pct') or 0.0), 2)
            ])

        president_obj = {
            'candidate': str(pw.get('Candidate', '')),
            'party': p_norm,
            'raw_party': str(pw.get('Party', '')),
            'votes': pw_votes,
            'pct': round(float(pw.get('Vote_Share_Pct') or 0.0), 2),
            'runner': str(prunner.get('Candidate', '')) if prunner else '',
            'runner_party': normalize_party(prunner.get('Party')) if prunner else '',
            'runner_votes': pr_votes,
            'runner_pct': round(float(prunner.get('Vote_Share_Pct') or 0.0), 2) if prunner else 0.0,
            'third': str(pthird.get('Candidate', '')) if pthird else '',
            'third_party': normalize_party(pthird.get('Party')) if pthird else '',
            'third_votes': int(pthird.get('Votes') or 0) if pthird else 0,
            'third_pct': round(float(pthird.get('Vote_Share_Pct') or 0.0), 2) if pthird else 0.0,
            'fourth': str(pfourth.get('Candidate', '')) if pfourth else '',
            'fourth_party': normalize_party(pfourth.get('Party')) if pfourth else '',
            'fourth_votes': int(pfourth.get('Votes') or 0) if pfourth else 0,
            'fourth_pct': round(float(pfourth.get('Vote_Share_Pct') or 0.0), 2) if pfourth else 0.0,
            'margin': pmargin,
            'cands': pres_cands
        }
        council_pres_tally[p_norm] = council_pres_tally.get(p_norm, 0) + 1

    # Councillor candidates grouped by (Ward, Seat)
    councillor_rows = [x for x in c_rows if str(x.get('Post_Type', '')).lower() != 'president']
    grouped = {}
    for cr in councillor_rows:
        w_no = str(cr.get('Ward', '') or '').strip()
        s_id = str(cr.get('Seat', '') or '').strip()
        if s_id == w_no:
            s_id = ''
        key = (w_no, s_id)
        if key not in grouped:
            grouped[key] = []
        grouped[key].append(cr)

    councillor_wards = []
    counc_tally = {}
    for (w_no, s_id), cand_rows in grouped.items():
        cand_rows.sort(key=lambda x: (x.get('Votes') or 0), reverse=True)
        winner = cand_rows[0]
        runner = cand_rows[1] if len(cand_rows) > 1 else None
        third = cand_rows[2] if len(cand_rows) > 2 else None
        fourth = cand_rows[3] if len(cand_rows) > 3 else None

        w_votes = int(winner.get('Votes') or 0)
        r_votes = int(runner.get('Votes') or 0) if runner else 0
        margin = w_votes - r_votes if runner else w_votes

        p_norm = normalize_party(winner.get('Party'))
        counc_tally[p_norm] = counc_tally.get(p_norm, 0) + 1
        council_tally_2025[p_norm] = council_tally_2025.get(p_norm, 0) + 1

        compact_cands = []
        for c in cand_rows:
            compact_cands.append([
                str(c.get('Candidate', '')),
                normalize_party(c.get('Party')),
                int(c.get('Votes') or 0),
                round(float(c.get('Vote_Share_Pct') or 0.0), 2)
            ])

        ward_lbl = f"{w_no}{s_id}" if s_id else str(w_no)

        councillor_wards.append({
            'ward': int(w_no) if w_no.isdigit() else w_no,
            'seat': s_id,
            'ward_label': ward_lbl,
            'candidate': str(winner.get('Candidate', '')),
            'party': p_norm,
            'votes': w_votes,
            'pct': round(float(winner.get('Vote_Share_Pct') or 0.0), 2),
            'runner': str(runner.get('Candidate', '')) if runner else '',
            'runner_party': normalize_party(runner.get('Party')) if runner else '',
            'runner_votes': r_votes,
            'runner_pct': round(float(runner.get('Vote_Share_Pct') or 0.0), 2) if runner else 0.0,
            'third': str(third.get('Candidate', '')) if third else '',
            'third_party': normalize_party(third.get('Party')) if third else '',
            'third_votes': int(third.get('Votes') or 0) if third else 0,
            'third_pct': round(float(third.get('Vote_Share_Pct') or 0.0), 2) if third else 0.0,
            'fourth': str(fourth.get('Candidate', '')) if fourth else '',
            'fourth_party': normalize_party(fourth.get('Party')) if fourth else '',
            'fourth_votes': int(fourth.get('Votes') or 0) if fourth else 0,
            'fourth_pct': round(float(fourth.get('Vote_Share_Pct') or 0.0), 2) if fourth else 0.0,
            'margin': margin,
            'cands': compact_cands
        })

    def c_sort(x):
        w = x['ward']
        s = x['seat']
        return (int(w) if str(w).isdigit() else 9999, str(w), str(s))
    councillor_wards.sort(key=c_sort)

    c_slug = re.sub(r'[^a-z0-9]+', '-', cname.lower()).strip('-')

    councils_list.append({
        'id': len(councils_list) + 1,
        'slug': c_slug,
        'name': cname,
        'raw_name': raw_name,
        'district': 'Maharashtra',
        'has_president': bool(president_obj),
        'president': president_obj,
        'seat_count': len(councillor_wards),
        'ward_count': len(set(x['ward'] for x in councillor_wards)),
        'party_tally': counc_tally,
        'winners': councillor_wards
    })

# Summary
comparison_parties = []
all_parties = sorted(list(set(list(corp_tally_2026.keys()) + list(corp_tally_2017.keys()))))
for p in all_parties:
    s26 = corp_tally_2026.get(p, 0)
    s17 = corp_tally_2017.get(p, 0)
    diff = s26 - s17
    comparison_parties.append({
        'party': p,
        'seats_2026': s26,
        'seats_2017': s17,
        'diff': diff
    })
comparison_parties.sort(key=lambda x: (x['seats_2026'] + x['seats_2017']), reverse=True)

summary_data = {
    'total_corps': len(corps_list),
    'total_councils': len(councils_list),
    'corps_seats_2026': sum(c['seat_count_2026'] for c in corps_list),
    'corps_seats_2017': sum(c['seat_count_2017'] for c in corps_list),
    'councils_seats_2025': sum(c['seat_count'] for c in councils_list),
    'corp_tally_2026': [{'party': p, 'seats': s} for p, s in sorted(corp_tally_2026.items(), key=lambda x: x[1], reverse=True)],
    'corp_tally_2017': [{'party': p, 'seats': s} for p, s in sorted(corp_tally_2017.items(), key=lambda x: x[1], reverse=True)],
    'council_tally_2025': [{'party': p, 'seats': s} for p, s in sorted(council_tally_2025.items(), key=lambda x: x[1], reverse=True)],
    'council_pres_tally': [{'party': p, 'seats': s} for p, s in sorted(council_pres_tally.items(), key=lambda x: x[1], reverse=True)],
    'comparison_parties': comparison_parties
}

print(f"Writing js/data-mh-local-bodies.js: {len(corps_list)} corps, {len(councils_list)} councils...")
with open('js/data-mh-local-bodies.js', 'w', encoding='utf-8') as f:
    f.write('window.MH_MUNC_SUMMARY = ' + json.dumps(summary_data, ensure_ascii=False) + ';\n')
    f.write('window.MH_CORPS = ' + json.dumps(corps_list, ensure_ascii=False) + ';\n')
    f.write('window.MH_COUNCILS = ' + json.dumps(councils_list, ensure_ascii=False) + ';\n')

print("js/data-mh-local-bodies.js updated successfully with all candidates!")
